import json
from collections import defaultdict
from pathlib import Path

import openpyxl

PROD_FILE = r"D:\Download\Production Plan Jul'26 Rev1..xlsx"
STOCK_FILE = r"D:\Download\MB52 (1).xlsx"
OUT = Path("outputs/sir_only_upload_data.json")


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def number(value):
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0


stock_wb = openpyxl.load_workbook(STOCK_FILE, read_only=True, data_only=True)
stock_ws = stock_wb["Data"]
stock_headers = [clean(v).lower() for v in next(stock_ws.iter_rows(min_row=1, max_row=1, values_only=True))]
stock_idx = {h: i for i, h in enumerate(stock_headers)}

components_by_id = {}
stock_by_id = defaultdict(float)
stock_qty_cols = [
    "unrestricted",
    "transit and transfer",
    "quality inspection",
    "restricted-use stock",
    "blocked stock",
    "returns",
]

for row in stock_ws.iter_rows(min_row=2, values_only=True):
    component_id = clean(row[stock_idx["material"]])
    if not component_id:
        continue
    name = clean(row[stock_idx["material description"]])
    category = clean(row[stock_idx["material group"]])
    uom = clean(row[stock_idx["base unit of measure"]])
    if component_id not in components_by_id:
        components_by_id[component_id] = {
            "Component ID": component_id,
            "Name": name,
            "Category": category,
            "UOM": uom,
        }
    stock_by_id[component_id] += sum(number(row[stock_idx[col]]) for col in stock_qty_cols if col in stock_idx)

components = sorted(components_by_id.values(), key=lambda x: x["Component ID"])
stock = [
    {
        "Component": components_by_id[component_id]["Name"],
        "Component ID": component_id,
        "Stock Value": round(qty, 6),
    }
    for component_id, qty in sorted(stock_by_id.items())
]

prod_wb = openpyxl.load_workbook(PROD_FILE, read_only=True, data_only=True)
prod_ws = prod_wb["FINAL WORKING."]
header_row = 4
prod_headers = [clean(v).lower() for v in next(prod_ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
prod_idx = {h: i for i, h in enumerate(prod_headers)}

fg_id_col = prod_idx["inel part no"]
desc_col = prod_idx["description"]
buyer_col = prod_idx["customer"]
qty_col = prod_idx["prodn. plan  jul'26 (qty"]

fg_by_id = {}
for row in prod_ws.iter_rows(min_row=header_row + 1, values_only=True):
    fg_id = clean(row[fg_id_col])
    description = clean(row[desc_col])
    qty = number(row[qty_col])
    if not fg_id or not description:
        continue
    if fg_id not in fg_by_id:
        fg_by_id[fg_id] = {
            "FG Part Number": fg_id,
            "Description": description,
            "Name": description,
            "Components": "",
            "Norms": "",
            "Buyer": clean(row[buyer_col]),
            "Production Quantity": 0,
        }
    fg_by_id[fg_id]["Production Quantity"] += qty

fg_products = [
    {
        "FG Part Number": item["FG Part Number"],
        "Description": item["Description"],
        "Name": item["Name"],
        "Components": "",
        "Norms": "",
        "Buyer": item["Buyer"],
    }
    for item in sorted(fg_by_id.values(), key=lambda x: x["FG Part Number"])
]

bom = [
    {
        "FG Part Number": item["FG Part Number"],
        "Name": item["Name"],
        "Production Quantity": int(item["Production Quantity"])
        if float(item["Production Quantity"]).is_integer()
        else round(item["Production Quantity"], 3),
    }
    for item in sorted(fg_by_id.values(), key=lambda x: x["FG Part Number"])
    if item["Production Quantity"] > 0
]

OUT.write_text(json.dumps({
    "components": components,
    "fg_products": fg_products,
    "bom": bom,
    "stock": stock,
}, indent=2), encoding="utf-8")

print(json.dumps({
    "components": len(components),
    "fgProducts": len(fg_products),
    "bom": len(bom),
    "stock": len(stock),
    "output": str(OUT),
}, indent=2))
