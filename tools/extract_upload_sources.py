import json
from collections import defaultdict
from pathlib import Path

import openpyxl

PROD_FILE = r"D:\Download\Production Plan Jul'26 Rev1..xlsx"
STOCK_FILE = r"D:\Download\MB52 (1).xlsx"
OUT = Path("outputs/upload_source_data.json")


def clean(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def num(value):
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return float(value)
    try:
        return float(str(value).replace(",", "").strip())
    except ValueError:
        return 0


stock_wb = openpyxl.load_workbook(STOCK_FILE, read_only=True, data_only=True)
stock_ws = stock_wb["Data"]
headers = [clean(v).lower() for v in next(stock_ws.iter_rows(min_row=1, max_row=1, values_only=True))]
idx = {h: i for i, h in enumerate(headers)}

stock_qty_cols = [
    "unrestricted",
    "transit and transfer",
    "quality inspection",
    "restricted-use stock",
    "blocked stock",
    "returns",
]

components_by_id = {}
stock_by_id = defaultdict(float)
for row in stock_ws.iter_rows(min_row=2, values_only=True):
    comp_id = clean(row[idx["material"]])
    if not comp_id:
        continue
    desc = clean(row[idx["material description"]])
    category = clean(row[idx["material group"]])
    uom = clean(row[idx["base unit of measure"]])
    if comp_id not in components_by_id:
        components_by_id[comp_id] = {
            "Component ID": comp_id,
            "Name": desc,
            "Category": category,
            "UOM": uom,
        }
    total_stock = sum(num(row[idx[col]]) for col in stock_qty_cols if col in idx)
    stock_by_id[comp_id] += total_stock

components = sorted(components_by_id.values(), key=lambda x: x["Component ID"])
stock = [
    {
        "Component": components_by_id[comp_id]["Name"],
        "Component ID": comp_id,
        "Stock Value": round(qty, 6),
    }
    for comp_id, qty in sorted(stock_by_id.items())
]

prod_wb = openpyxl.load_workbook(PROD_FILE, read_only=True, data_only=True)
prod_ws = prod_wb["FINAL WORKING."]
header_row = 4
prod_headers = [clean(v).lower() for v in next(prod_ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))]
prod_idx = {h: i for i, h in enumerate(prod_headers)}

fg_id_col = prod_idx["inel part no"]
desc_col = prod_idx["description"]
customer_col = prod_idx["customer"]
qty_col = prod_idx["prodn. plan  jul'26 (qty"]

fg_by_id = {}
for row in prod_ws.iter_rows(min_row=header_row + 1, values_only=True):
    fg_id = clean(row[fg_id_col])
    desc = clean(row[desc_col])
    qty = num(row[qty_col])
    if not fg_id or not desc or qty <= 0:
        continue
    customer = clean(row[customer_col])
    if fg_id not in fg_by_id:
        fg_by_id[fg_id] = {
            "FG Part Number": fg_id,
            "Name": desc,
            "Description": desc,
            "Buyer": customer,
            "Production Quantity": 0,
        }
    fg_by_id[fg_id]["Production Quantity"] += qty

fgs = sorted(fg_by_id.values(), key=lambda x: x["FG Part Number"])
bom = [
    {
        "FG Part Number": item["FG Part Number"],
        "Name": item["Name"],
        "Production Quantity": int(item["Production Quantity"])
        if float(item["Production Quantity"]).is_integer()
        else round(item["Production Quantity"], 3),
    }
    for item in fgs
]

# These two source files do not contain FG -> component norms. For upload testing,
# attach real component IDs to each real FG in a deterministic small pattern.
component_ids = [item["Component ID"] for item in components if item["Component ID"]]
norm_pattern = [1, 2, 0.5, 3]
fg_products = []
for fg_pos, fg in enumerate(fgs):
    for offset in range(4):
        comp_id = component_ids[(fg_pos * 7 + offset * 13) % len(component_ids)]
        fg_products.append({
            "FG Part Number": fg["FG Part Number"],
            "Description": fg["Description"],
            "Name": fg["Name"],
            "Component ID": comp_id,
            "Norms": norm_pattern[(fg_pos + offset) % len(norm_pattern)],
            "Buyer": fg["Buyer"],
        })

OUT.parent.mkdir(parents=True, exist_ok=True)
OUT.write_text(json.dumps({
    "components": components,
    "fg_products": fg_products,
    "bom": bom,
    "stock": stock,
    "notes": {
        "production_source": PROD_FILE,
        "stock_source": STOCK_FILE,
        "fg_component_mapping": "Generated for upload testing because source files do not include FG-to-component norms.",
    },
}, indent=2), encoding="utf-8")

print(json.dumps({
    "components": len(components),
    "fg_rows": len(fg_products),
    "fg_unique": len(fgs),
    "bom_rows": len(bom),
    "stock_rows": len(stock),
    "output": str(OUT),
}, indent=2))
