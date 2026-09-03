import openpyxl

files = [
    r"D:\Download\Production Plan Jul'26 Rev1..xlsx",
    r"D:\Download\MB52 (1).xlsx",
]

for path in files:
    print(f"\nFILE {path}")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print("sheets", wb.sheetnames)
    for ws in wb.worksheets:
        print("SHEET", ws.title, "max", ws.max_row, ws.max_column)
        shown = 0
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = list(row[:28])
            if any(v is not None for v in vals):
                print(row_idx, vals)
                shown += 1
            if shown >= 12:
                break
        print("---")
